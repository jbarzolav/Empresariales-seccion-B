"""
Script to generate a highly polished, professional PowerPoint presentation on 'Contaminacion vehicular urbana en Lima'
incorporating a formal academic cover slide (Slide 1) as requested, followed by topic intro, justification,
2x2 grid problem statement with image placeholders, and interview questions.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import shutil
import pathlib

def create_presentation():
    prs = Presentation()
    # Set 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank layout

    # Professional Design Palette (Tech Blues & Cyan on Clean White)
    BG_COLOR = RGBColor(255, 255, 255)       # Clean white background
    PRIMARY_COLOR = RGBColor(15, 23, 42)     # Dark graphite text
    ACCENT_COLOR = RGBColor(14, 165, 233)    # Tech Cyan / Blue accent
    ACCENT_GREEN = RGBColor(16, 185, 129)    # Emerald Green accent
    CARD_BG = RGBColor(248, 250, 252)        # Very light gray for cards
    TEXT_MUTED = RGBColor(71, 85, 105)       # Slate gray for descriptions
    CARD_BORDER = RGBColor(226, 232, 240)    # Subtle border

    def set_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_COLOR

    # ==========================================
    # SLIDE 1: Carátula Académica Formal
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_background(slide1)

    # Accent top bar (Tech Cyan)
    top_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(0.8), Inches(2.5), Inches(0.12))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = ACCENT_COLOR
    top_bar.line.fill.background()

    # Main Card container for Cover
    card1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1.2), Inches(11.333), Inches(5.6))
    card1.fill.solid()
    card1.fill.fore_color.rgb = CARD_BG
    card1.line.color.rgb = CARD_BORDER

    tf1 = card1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.8)
    tf1.margin_right = Inches(0.8)
    tf1.margin_top = Inches(0.5)

    # Title Principal
    p1 = tf1.paragraphs[0]
    p1.text = "INVESTIGACION E INNOVACION TECNOLOGICA"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = PRIMARY_COLOR
    p1.space_after = Pt(8)

    # Subtítulo (Sección)
    p2 = tf1.add_paragraph()
    p2.text = "SECCION C24-B"
    p2.font.size = Pt(16)
    p2.font.bold = True
    p2.font.color.rgb = ACCENT_COLOR
    p2.space_after = Pt(16)

    # Carrera
    p3 = tf1.add_paragraph()
    p3.text = "CARRERA DE DISEÑO Y DESARROLLO DE SOFTWARE"
    p3.font.size = Pt(14)
    p3.font.bold = True
    p3.font.color.rgb = TEXT_MUTED
    p3.space_after = Pt(16)

    # Docente
    p4 = tf1.add_paragraph()
    p4.text = "DOCENTE: [Nombre del Profesor]"
    p4.font.size = Pt(13)
    p4.font.bold = True
    p4.font.color.rgb = PRIMARY_COLOR
    p4.space_after = Pt(16)

    # Integrantes title
    p5 = tf1.add_paragraph()
    p5.text = "INTEGRANTES:"
    p5.font.size = Pt(12)
    p5.font.bold = True
    p5.font.color.rgb = ACCENT_GREEN
    p5.space_after = Pt(4)

    # Integrantes list (1 to 5)
    integrantes = [
        "[Nombre Apellido]",
        "[Nombre Apellido]",
        "[Nombre Apellido]",
        "[Nombre Apellido]",
        "[Nombre Apellido]"
    ]
    for integ in integrantes:
        pi = tf1.add_paragraph()
        pi.text = f"• {integ}"
        pi.font.size = Pt(12)
        pi.font.color.rgb = TEXT_MUTED
        pi.space_after = Pt(2)

    # Footer year/semester inside slide bottom
    footer_box = slide1.shapes.add_textbox(Inches(1), Inches(6.8), Inches(11.333), Inches(0.5))
    tf_f = footer_box.text_frame
    pf = tf_f.paragraphs[0]
    pf.text = "2026-2"
    pf.font.size = Pt(12)
    pf.font.bold = True
    pf.font.color.rgb = TEXT_MUTED
    pf.alignment = PP_ALIGN.RIGHT

    notes1 = slide1.notes_slide.notes_text_frame
    notes1.text = "Transición: Desvanecer (Fade) o Transformación (Morph). Portada formal institucional."

    # ==========================================
    # SLIDE 2: Presentación del Tema (Contaminación vehicular en Lima)
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_background(slide2)

    t_box2 = slide2.shapes.add_textbox(Inches(1), Inches(0.8), Inches(11.333), Inches(1))
    tf2_t = t_box2.text_frame
    p_t2 = tf2_t.paragraphs[0]
    p_t2.text = "Contaminación vehicular urbana en Lima"
    p_t2.font.size = Pt(32)
    p_t2.font.bold = True
    p_t2.font.color.rgb = PRIMARY_COLOR

    card2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2.0), Inches(11.333), Inches(4.5))
    card2.fill.solid()
    card2.fill.fore_color.rgb = CARD_BG
    card2.line.color.rgb = CARD_BORDER

    left_accent2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2.0), Inches(0.15), Inches(4.5))
    left_accent2.fill.solid()
    left_accent2.fill.fore_color.rgb = ACCENT_COLOR
    left_accent2.line.fill.background()

    tf2 = card2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = Inches(1.2)
    tf2.margin_right = Inches(1.0)
    tf2.margin_top = Inches(1.2)

    p_sub2 = tf2.paragraphs[0]
    p_sub2.text = "Elección del tema:"
    p_sub2.font.size = Pt(16)
    p_sub2.font.bold = True
    p_sub2.font.color.rgb = ACCENT_GREEN
    p_sub2.space_after = Pt(8)

    p_body2 = tf2.add_paragraph()
    p_body2.text = "Se eligió porque el impacto del tráfico y la mala calidad del aire afectan directamente la salud respiratoria y la vida cotidiana de las personas que viven en zonas urbanas densas de Lima."
    p_body2.font.size = Pt(20)
    p_body2.font.color.rgb = PRIMARY_COLOR
    p_body2.line_spacing = 1.3

    notes2 = slide2.notes_slide.notes_text_frame
    notes2.text = "Transición: Transformación (Morph) o Desvanecer (Fade). Animación de aparición gradual."

    # ==========================================
    # SLIDE 3: Justificación (Importancia y Pertinencia)
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_background(slide3)

    t_box3 = slide3.shapes.add_textbox(Inches(1), Inches(0.8), Inches(11.333), Inches(1))
    tf3_t = t_box3.text_frame
    p_t3 = tf3_t.paragraphs[0]
    p_t3.text = "Importancia y Pertinencia"
    p_t3.font.size = Pt(32)
    p_t3.font.bold = True
    p_t3.font.color.rgb = PRIMARY_COLOR

    card3 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2.0), Inches(11.333), Inches(4.5))
    card3.fill.solid()
    card3.fill.fore_color.rgb = CARD_BG
    card3.line.color.rgb = CARD_BORDER

    left_accent3 = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2.0), Inches(0.15), Inches(4.5))
    left_accent3.fill.solid()
    left_accent3.fill.fore_color.rgb = ACCENT_GREEN
    left_accent3.line.fill.background()

    tf3 = card3.text_frame
    tf3.word_wrap = True
    tf3.margin_left = Inches(1.2)
    tf3.margin_right = Inches(1.0)
    tf3.margin_top = Inches(1.2)

    p_body3 = tf3.paragraphs[0]
    p_body3.text = "La contaminación del aire es un problema crítico de salud pública en las ciudades actuales. Analizar este impacto es pertinente para concienciar y buscar alternativas que protejan el bienestar de los ciudadanos frente a la exposición diaria a gases contaminantes."
    p_body3.font.size = Pt(22)
    p_body3.font.color.rgb = PRIMARY_COLOR
    p_body3.line_spacing = 1.4

    notes3 = slide3.notes_slide.notes_text_frame
    notes3.text = "Transición: Transformación (Morph) o Desvanecer (Fade). Animación de la tarjeta con efecto de elevación."

    # ==========================================
    # SLIDE 4: Planteamiento del Problema (2x2 Grid with image placeholders)
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_background(slide4)

    t_box4 = slide4.shapes.add_textbox(Inches(1), Inches(0.6), Inches(11.333), Inches(0.8))
    tf4_t = t_box4.text_frame
    p_t4 = tf4_t.paragraphs[0]
    p_t4.text = "Planteamiento del Problema"
    p_t4.font.size = Pt(32)
    p_t4.font.bold = True
    p_t4.font.color.rgb = PRIMARY_COLOR

    blocks = [
        ("Problema Central", "Falta de información accesible sobre los niveles reales de contaminación exterior al momento de ventilar los hogares en zonas de alto tráfico.", "[ PRO ]"),
        ("Causas", "Alta congestión vehicular, antigüedad del parque automotor y emisiones directas de material particulado (PM2.5 y PM10) en las vías urbanas.", "[ CAU ]"),
        ("Consecuencias", "Exposición involuntaria a gases nocivos, aumento de problemas respiratorios e incertidumbre diaria sobre cuándo es seguro abrir las ventanas.", "[ CON ]"),
        ("Afectados", "Vecinos y familias que residen en viviendas ubicadas frente a avenidas con alto flujo vehicular en Lima.", "[ AFE ]")
    ]

    grid_coords = [
        (Inches(1), Inches(1.6)),     # Top-Left
        (Inches(6.8), Inches(1.6)),   # Top-Right
        (Inches(1), Inches(4.5)),     # Bottom-Left
        (Inches(6.8), Inches(4.5))    # Bottom-Right
    ]

    for i, (title, desc, tag) in enumerate(blocks):
        x, y = grid_coords[i]
        
        box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.5), Inches(2.5))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = CARD_BORDER

        # Image placeholder box on the right side of each card
        img_box = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, x + Inches(3.6), y + Inches(0.35), Inches(1.5), Inches(1.8))
        img_box.fill.solid()
        img_box.fill.fore_color.rgb = RGBColor(238, 242, 246)
        img_box.line.color.rgb = RGBColor(203, 213, 225)
        tf_img = img_box.text_frame
        tf_img.word_wrap = True
        p_img = tf_img.paragraphs[0]
        p_img.text = "[ Imagen / Icono ]"
        p_img.font.size = Pt(10)
        p_img.font.color.rgb = RGBColor(148, 163, 184)
        p_img.alignment = PP_ALIGN.CENTER

        badge = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.3), y + Inches(0.25), Inches(0.8), Inches(0.3))
        badge.fill.solid()
        badge.fill.fore_color.rgb = ACCENT_COLOR if i % 2 == 0 else ACCENT_GREEN
        badge.line.fill.background()
        tf_bg = badge.text_frame
        p_bg = tf_bg.paragraphs[0]
        p_bg.text = tag
        p_bg.font.size = Pt(10)
        p_bg.font.bold = True
        p_bg.font.color.rgb = RGBColor(255, 255, 255)
        p_bg.alignment = PP_ALIGN.CENTER

        tf_b = box.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = Inches(0.3)
        tf_b.margin_right = Inches(2.0)
        tf_b.margin_top = Inches(0.65)

        pb_t = tf_b.paragraphs[0]
        pb_t.text = title
        pb_t.font.size = Pt(16)
        pb_t.font.bold = True
        pb_t.font.color.rgb = PRIMARY_COLOR
        pb_t.space_after = Pt(4)

        pb_d = tf_b.add_paragraph()
        pb_d.text = desc
        pb_d.font.size = Pt(12)
        pb_d.font.color.rgb = TEXT_MUTED
        pb_d.line_spacing = 1.2

    notes4 = slide4.notes_slide.notes_text_frame
    notes4.text = "Transición: Transformación (Morph) o Desvanecer (Fade). Animación secuencial (en cascada) para los 4 bloques en cuadrícula."

    # ==========================================
    # SLIDE 5: Guion de Entrevistas (Parte 1: Q1-Q5)
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_background(slide5)

    t_box5 = slide5.shapes.add_textbox(Inches(1), Inches(0.6), Inches(11.333), Inches(0.8))
    tf5_t = t_box5.text_frame
    p_t5 = tf5_t.paragraphs[0]
    p_t5.text = "Guion de Preguntas para Entrevistas (1/2)"
    p_t5.font.size = Pt(32)
    p_t5.font.bold = True
    p_t5.font.color.rgb = PRIMARY_COLOR

    q_card1 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1.5), Inches(11.333), Inches(5.4))
    q_card1.fill.solid()
    q_card1.fill.fore_color.rgb = CARD_BG
    q_card1.line.color.rgb = CARD_BORDER

    tf_q1 = q_card1.text_frame
    tf_q1.word_wrap = True
    tf_q1.margin_left = Inches(0.8)
    tf_q1.margin_right = Inches(0.8)
    tf_q1.margin_top = Inches(0.5)

    questions_part1 = [
        "¿Con qué frecuencia nota la presencia de tráfico vehicular pesado o acumulación de humo frente a su domicilio?",
        "¿Considera que el aire que se respira en su zona residencial es saludable para su día a día?",
        "¿Tiene acceso a información clara o reportes sencillos sobre el nivel de contaminación en su distrito?",
        "¿Qué tan fácil o difícil le resulta interpretar los términos técnicos sobre contaminación del aire?",
        "¿En qué momentos del día suele abrir las ventanas de su vivienda para ventilar los espacios?"
    ]

    for idx, q in enumerate(questions_part1, 1):
        pq = tf_q1.paragraphs[0] if idx == 1 else tf_q1.add_paragraph()
        pq.text = f"{idx}. {q}"
        pq.font.size = Pt(17)
        pq.font.color.rgb = PRIMARY_COLOR
        pq.space_after = Pt(16)

    notes5 = slide5.notes_slide.notes_text_frame
    notes5.text = "Transición: Desvanecer (Fade). Animación de aparición gradual (Fade) secuencial ítem por ítem."

    # ==========================================
    # SLIDE 6: Guion de Entrevistas (Parte 2: Q6-Q10)
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_background(slide6)

    t_box6 = slide6.shapes.add_textbox(Inches(1), Inches(0.6), Inches(11.333), Inches(0.8))
    tf6_t = t_box6.text_frame
    p_t6 = tf6_t.paragraphs[0]
    p_t6.text = "Guion de Preguntas para Entrevistas (2/2)"
    p_t6.font.size = Pt(32)
    p_t6.font.bold = True
    p_t6.font.color.rgb = PRIMARY_COLOR

    q_card2 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1.5), Inches(11.333), Inches(5.4))
    q_card2.fill.solid()
    q_card2.fill.fore_color.rgb = CARD_BG
    q_card2.line.color.rgb = CARD_BORDER

    tf_q2 = q_card2.text_frame
    tf_q2.word_wrap = True
    tf_q2.margin_left = Inches(0.8)
    tf_q2.margin_right = Inches(0.8)
    tf_q2.margin_top = Inches(0.5)

    questions_part2 = [
        "¿Qué factores o señales toma en cuenta para decidir si es un buen momento para abrir las ventanas?",
        "¿Alguna vez ha tenido que mantener las ventanas cerradas todo el día debido al humo o la congestión exterior?",
        "¿De qué manera la contaminación vehicular exterior afecta la comodidad o la tranquilidad en su hogar?",
        "¿Le gustaría recibir avisos directos sobre cuándo el aire de su calle está demasiado contaminado?",
        "¿Qué tipo de herramientas o indicaciones le gustaría tener a la mano para saber cuándo es seguro ventilar su casa?"
    ]

    for idx, q in enumerate(questions_part2, 6):
        pq = tf_q2.paragraphs[0] if idx == 6 else tf_q2.add_paragraph()
        pq.text = f"{idx}. {q}"
        pq.font.size = Pt(17)
        pq.font.color.rgb = PRIMARY_COLOR
        pq.space_after = Pt(16)

    notes6 = slide6.notes_slide.notes_text_frame
    notes6.text = "Transición: Desvanecer (Fade). Animación de aparición gradual (Fade) secuencial ítem por ítem."

    # Save presentation directly to Downloads folder
    downloads = pathlib.Path(r'C:\Users\josem\Downloads')
    downloads.mkdir(exist_ok=True)
    dest_path = downloads / "Contaminacion_Vehicular_Lima.pptx"
    try:
        if dest_path.exists():
            dest_path.unlink()
        prs.save(dest_path)
        print(f"Presentation successfully created and saved to: {dest_path}")
    except PermissionError:
        alt_path = downloads / "Contaminacion_Vehicular_Lima_actualizado.pptx"
        prs.save(alt_path)
        print(f"Original file is open in PowerPoint. Saved as: {alt_path}")

if __name__ == "__main__":
    create_presentation()
